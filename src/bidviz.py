#!/usr/bin/env python3
"""施工组织设计可视化：事实提取、工程生图、A3 PDF 与审计。"""

from __future__ import annotations

import argparse
import base64
import hashlib
import importlib.util
import json
import os
import re
import shutil
import sys
import textwrap
import urllib.error
import urllib.request
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import landscape, A3
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

try:
    import fitz
except ImportError:  # pragma: no cover
    fitz = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


A3_PX = (4961, 3508)
A3_MM = (420, 297)
MAX_REFERENCES = 4
STATUS = {"exact", "normalized", "derived", "suggested", "ambiguous", "conflict", "approved"}
BLOCKING_STATUS = {"suggested", "ambiguous", "conflict"}
OCR_AVAILABLE = importlib.util.find_spec("pytesseract") is not None and shutil.which("tesseract") is not None
VISUAL_LABELS = {
    "site-reality": "项目现场实景",
    "drawing-constrained": "基于项目图纸的施工示意图",
    "reference-constrained": "基于项目资料的效果示意图",
    "generic": "AI通用施工示意图",
}


@dataclass
class Source:
    source_id: str
    path: str
    kind: str
    sha256: str
    title: str = ""


@dataclass
class Claim:
    claim_id: str
    work_item_id: str
    section_path: str
    text: str
    raw_quote: str
    source_id: str
    locator: str
    claim_type: str
    status: str = "exact"
    normalized_value: str | None = None
    unit: str | None = None

    def validate(self) -> None:
        if self.status not in STATUS:
            raise ValueError(f"未知事实状态：{self.status}")
        if not all((self.claim_id, self.text, self.raw_quote, self.source_id, self.locator)):
            raise ValueError(f"事实缺少溯源字段：{self.claim_id or '<unknown>'}")


@dataclass
class Issue:
    issue_id: str
    claim_ids: list[str]
    severity: str
    category: str
    message: str
    blocking: bool = True
    resolved: bool = False


@dataclass
class PageSpec:
    page_id: str
    work_item_id: str
    title: str
    visual_type: str
    claim_ids: list[str]
    reference_assets: list[str] = field(default_factory=list)
    truth_level: str = "generic"
    scoring_points: list[str] = field(default_factory=list)
    prompt: str = ""


@dataclass
class Project:
    name: str
    sources: list[Source] = field(default_factory=list)
    claims: list[Claim] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    pages: list[PageSpec] = field(default_factory=list)


class BidVizError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def infer_kind(path: Path) -> str:
    suffix = path.suffix.lower()
    return {
        ".docx": "construction-document",
        ".pdf": "pdf",
        ".xlsx": "spreadsheet",
        ".xlsm": "spreadsheet",
        ".png": "reference-image",
        ".jpg": "reference-image",
        ".jpeg": "reference-image",
        ".webp": "reference-image",
        ".txt": "text",
        ".md": "text",
    }.get(suffix, "other")


def build_manifest(input_dir: Path) -> dict[str, Any]:
    files = sorted(p for p in input_dir.rglob("*") if p.is_file())
    sources = []
    counters: dict[str, int] = {}
    prefixes = {
        "construction-document": "DOC",
        "pdf": "PDF",
        "spreadsheet": "XLSX",
        "reference-image": "IMG",
        "text": "TXT",
        "other": "FILE",
    }
    for path in files:
        kind = infer_kind(path)
        prefix = prefixes[kind]
        counters[prefix] = counters.get(prefix, 0) + 1
        sources.append(asdict(Source(
            source_id=f"{prefix}-{counters[prefix]:03d}",
            path=path.relative_to(input_dir).as_posix(),
            kind=kind,
            sha256=sha256_file(path),
            title=path.stem,
        )))
    return {"version": 1, "created_at": utc_now(), "input_root": str(input_dir), "sources": sources}


def save_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def extract_docx(path: Path) -> list[str]:
    if Document is None:
        raise BidVizError("缺少 python-docx")
    doc = Document(path)
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.strip(" |"): lines.append(line)
    return lines


def extract_pdf(path: Path) -> list[str]:
    if fitz is None:
        raise BidVizError("缺少 PyMuPDF")
    doc = fitz.open(path)
    return [line.strip() for page in doc for line in page.get_text().splitlines() if line.strip()]


def extract_xlsx(path: Path) -> list[str]:
    if load_workbook is None:
        raise BidVizError("缺少 openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values: lines.append(f"[{ws.title}] " + " | ".join(values))
    return lines


def extract_lines(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".docx": return extract_docx(path)
    if suffix == ".pdf": return extract_pdf(path)
    if suffix in {".xlsx", ".xlsm"}: return extract_xlsx(path)
    if suffix in {".txt", ".md"}: return [x.strip() for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
    return []


def is_heading(line: str) -> bool:
    return bool(re.match(r"^(第[一二三四五六七八九十百\d]+[章节]|[一二三四五六七八九十]+、|\([一二三四五六七八九十]+\)|（[一二三四五六七八九十]+）)", line)) and len(line) < 80


def claim_type(text: str) -> str:
    if re.search(r"严禁|不得|禁止", text) and re.search(r"安全|危险|停工|基坑|基槽|边坡|车辆|人员|夜间|警示", text): return "PROHIBITION"
    if re.search(r"验收|检查|报验|确认|记录", text): return "INSPECTION"
    if re.search(r"危险|塌方|裂缝|应急|警示|安全", text): return "RISK_CONTROL"
    if re.search(r"如果|如遇|当.+时|超过|不足", text): return "CONDITION"
    if re.search(r"\d", text): return "PARAMETER"
    if re.search(r"严禁|不得|禁止", text): return "QUALITY_CONTROL"
    return "ACTION"


def extract_claims(lines: Iterable[str], source_id: str, work_item_id: str = "AUTO") -> list[Claim]:
    section: list[str] = []
    claims: list[Claim] = []
    for idx, raw in enumerate(lines, 1):
        line = re.sub(r"\s+", " ", raw).strip()
        if not line: continue
        if is_heading(line):
            level_hint = 1 if line.startswith("第") else 2 if re.match(r"^[一二三四五六七八九十]+、", line) else 3
            section = section[:level_hint - 1] + [line]
            continue
        parts = [p.strip() for p in re.split(r"(?<=[。；;])", line) if p.strip()]
        for part in parts:
            cid = f"C-{len(claims)+1:05d}"
            status = "ambiguous" if re.search(r"\d+、\d+\s*[~～至]\s*\d+、\d+", part) else "exact"
            claims.append(Claim(
                claim_id=cid,
                work_item_id=work_item_id,
                section_path=" / ".join(section) or "未识别章节",
                text=part,
                raw_quote=part,
                source_id=source_id,
                locator=f"line:{idx}",
                claim_type=claim_type(part),
                status=status,
            ))
    return claims


def detect_issues(claims: list[Claim]) -> list[Issue]:
    issues: list[Issue] = []
    for claim in claims:
        text = claim.text
        if claim.status == "ambiguous":
            issues.append(Issue(f"I-{len(issues)+1:04d}", [claim.claim_id], "high", "ambiguous-number", f"疑似小数点 OCR 错误：{text}"))
        if re.search(r"不少于\s*\d+\s*[。；;]?$", text):
            issues.append(Issue(f"I-{len(issues)+1:04d}", [claim.claim_id], "high", "missing-unit", f"数值缺少单位：{text}"))
        if re.search(r"\b(?:A|O)\d{2,3}\b", text):
            issues.append(Issue(f"I-{len(issues)+1:04d}", [claim.claim_id], "medium", "suspected-ocr", f"疑似符号 OCR 错误：{text}"))
    # 同一工作项中“上卷”高度冲突。
    heights: dict[str, list[tuple[str, str]]] = {}
    for claim in claims:
        if "上卷" in claim.text:
            for n in re.findall(r"(\d+)\s*mm", claim.text, re.I):
                heights.setdefault(claim.work_item_id, []).append((n, claim.claim_id))
    for wid, values in heights.items():
        unique = {v for v, _ in values}
        if len(unique) > 1:
            ids = [cid for _, cid in values]
            issues.append(Issue(f"I-{len(issues)+1:04d}", ids, "high", "conflict", f"{wid} 的防水上卷高度存在冲突：{', '.join(sorted(unique))}mm"))
    return issues


def visual_plan(work_item_id: str, title: str, claims: list[Claim], references: list[str]) -> list[PageSpec]:
    joined = "\n".join(c.text for c in claims)
    pages: list[PageSpec] = []
    def add(kind: str, page_title: str, selected: list[Claim]) -> None:
        idx = len(pages) + 1
        truth = "drawing-constrained" if references else "generic"
        if kind == "location-context" and not references:
            page_title = f"{title}通用空间关系示意"
        pages.append(PageSpec(
            page_id=f"{work_item_id}-{idx:02d}", work_item_id=work_item_id,
            title=page_title, visual_type=kind,
            claim_ids=[c.claim_id for c in selected[:18]], reference_assets=references[:MAX_REFERENCES], truth_level=truth,
        ))
    spatial = bool(references) or bool(re.search(r"位置|走向|道路|房屋|建筑|周边|线路|运输|出入口|两侧|邻近|交叉", joined))
    if spatial:
        add("location-context", f"{title}项目位置与周边关系", [c for c in claims if c.claim_type in {"ACTION", "CONDITION", "PARAMETER"}])
    add("construction-scene", f"{title}施工场景示意", claims)
    if re.search(r"坡|厚|深|宽|高|间距|断面|分层|垫层|沟", joined):
        add("section-or-exploded", f"{title}构造剖面与分层", [c for c in claims if c.claim_type == "PARAMETER"])
    if sum(c.claim_type == "PARAMETER" for c in claims) >= 3:
        add("parameter-control", f"{title}关键参数控制", [c for c in claims if c.claim_type in {"PARAMETER", "INSPECTION"}])
    finished_product = bool(re.search(r"成品|盖板安装|顺直|平整|色泽一致|无空鼓|无龟裂|无积水|外观|压光后", joined))
    has_risk = any(c.claim_type in {"RISK_CONTROL", "PROHIBITION"} for c in claims)
    if has_risk:
        add("quality-safety", f"{title}质量安全与风险处置", [c for c in claims if c.claim_type in {"RISK_CONTROL", "PROHIBITION", "INSPECTION", "CONDITION"}])
    elif finished_product:
        add("finished-product", f"{title}完工实物效果示意", [c for c in claims if re.search(r"成品|安装|顺直|平整|色泽|空鼓|龟裂|积水|外观|压光", c.text)])
    return pages[:5]


def compile_prompt(page: PageSpec, claims_by_id: dict[str, Claim]) -> str:
    facts = [claims_by_id[cid].text for cid in page.claim_ids if cid in claims_by_id]
    types = {
        "location-context": "工程项目位置与周边关系鸟瞰图",
        "construction-scene": "专业施工技术轴测场景图",
        "section-or-exploded": "工程构造剖面或分层爆炸图",
        "parameter-control": "留有大面积表格区域的工程主题背景图",
        "quality-safety": "带风险空间关系的施工安全示意图",
        "finished-product": "符合验收观感要求的工程完工实物效果示意图",
    }
    truth_instruction = "严格保持参考图中的真实线路、构件位置和主要几何关系。" if page.reference_assets else "无项目平面图参考；只生成通用空间关系，不得暗示真实项目坐标、建筑布局或现场位置。"
    return "\n".join([
        f"生成{types.get(page.visual_type, '专业工程技术示意图')}。",
        truth_instruction,
        "需要表现的事实：" + "；".join(facts[:10]),
        "风格：严谨、克制、工程技术插画、A3横向构图、浅灰白背景、蓝灰主色、危险区域少量橙红。",
        "绝对禁止生成任何中文、英文、数字、尺寸、表格文字、标题栏、水印、Logo、空白表格或工程图框；只输出纯视觉工程插图。",
        "禁止增加事实中没有的设备、尺寸、人员职责和项目位置。",
        "画面只保留主体与自然留白，不要绘制页面模板或表格区域；标题、控制表和免责声明由后续程序叠加。",
    ])


def pideck_image_config(provider_name: str | None = None, model: str | None = None) -> dict[str, str]:
    """兼容旧调用；凭据实际由独立配置层解析，PiDeck 只是最后回退。"""
    try:
        from bidviz_config import get, provider_credentials
        name = provider_name or get().provider
        cfg = provider_credentials(name)
        return {"base_url": cfg.get("baseUrl", ""),
                "api_key": cfg.get("apiKey", ""),
                "model": model or get().model}
    except RuntimeError as exc:
        raise BidVizError(str(exc)) from None


class ImageAPI:
    def __init__(self, base_url: str | None = None, api_key: str | None = None, model: str | None = None, provider_name: str | None = None):
        env_base = os.getenv("BIDVIZ_BASE_URL") or os.getenv("CUSTOM_IMAGE_API_BASE")
        env_key = os.getenv("BIDVIZ_API_KEY") or os.getenv("CUSTOM_IMAGE_API_KEY")
        env_model = os.getenv("BIDVIZ_MODEL") or os.getenv("CUSTOM_IMAGE_MODEL")
        configured = bool(base_url or api_key or env_base or env_key)
        cfg = {} if configured else pideck_image_config(
            provider_name or os.getenv("BIDVIZ_PROVIDER") or os.getenv("CUSTOM_IMAGE_PROVIDER"), model)
        self.base_url = (base_url or env_base or cfg.get("base_url", "")).rstrip("/")
        self.api_key = api_key or env_key or cfg.get("api_key", "")
        self.model = model or env_model or cfg.get("model", "gpt-image-2")
        if not self.base_url or not self.api_key:
            raise BidVizError("缺少生图凭据：请设置 BIDVIZ_BASE_URL / BIDVIZ_API_KEY")

    @property
    def endpoint(self) -> str:
        """按根地址推导生图端点：已是端点则直接用；已带版本段则追加；裸域名补 /v1。"""
        base = self.base_url.rstrip("/")
        if base.endswith("/images/generations"):
            return base
        if re.search(r"/(v\d+[a-z]*|api|api/v\d+)$", base):
            return base + "/images/generations"
        return base + "/v1/images/generations"

    def models(self) -> list[str]:
        base = self.base_url[:-len("/images/generations")] if self.base_url.endswith("/images/generations") else self.base_url
        req = urllib.request.Request(base.rstrip("/") + "/models", headers={"Authorization": f"Bearer {self.api_key}"})
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                payload = json.load(resp)
        except urllib.error.HTTPError as exc:
            raise BidVizError(f"模型探测 HTTP {exc.code}") from None
        return [item.get("id", "") for item in payload.get("data", []) if isinstance(item, dict) and item.get("id")]

    @staticmethod
    def _data_uri(path: Path) -> str:
        mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}.get(path.suffix.lower(), "image/png")
        return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii")

    def generate(self, prompt: str, output: Path, references: list[Path] | None = None,
                 recorder: dict[str, Any] | None = None, size: str | None = None) -> Path:
        refs = references or []
        if len(refs) > MAX_REFERENCES:
            raise BidVizError(f"参考图最多 {MAX_REFERENCES} 张")
        body: dict[str, Any] = {"model": self.model, "prompt": prompt, "n": 1, "response_format": "b64_json"}
        if refs: body["image"] = [self._data_uri(p) for p in refs]
        if size: body["size"] = size
        if recorder is not None:
            recorder.clear()
            recorder.update({"endpoint": self.endpoint, "model": self.model, "reference_count": len(refs), "prompt_sha256": hashlib.sha256(prompt.encode("utf-8")).hexdigest()})
        req = urllib.request.Request(
            self.endpoint,
            data=json.dumps(body).encode("utf-8"),
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(req, timeout=240) as resp:
                data = json.load(resp)
        except urllib.error.HTTPError as exc:
            raw = exc.read().decode("utf-8", "replace")
            try:
                parsed = json.loads(raw)
                error = parsed.get("error", parsed)
                message = error.get("message", str(error)) if isinstance(error, dict) else str(error)
            except Exception:
                message = raw
            raise BidVizError(f"生图 API HTTP {exc.code}: {message[:500]}") from None
        items = data.get("data") or data.get("images") or []
        if not items:
            raise BidVizError("生图 API 未返回图片")
        item = items[0]
        if item.get("b64_json"):
            content = base64.b64decode(item["b64_json"])
        elif item.get("url"):
            with urllib.request.urlopen(item["url"], timeout=120) as resp:
                content = resp.read()
        else:
            raise BidVizError("响应没有 b64_json 或 url")
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(content)
        if recorder is not None:
            recorder.update({"output": str(output), "output_sha256": hashlib.sha256(content).hexdigest(), "generated_at": utc_now()})
        return output


def font_path(bold: bool = False) -> str:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists(): return candidate
    raise BidVizError("未找到可用中文字体")


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path(bold), size)


def placeholder_visual(page: PageSpec) -> Image.Image:
    image = Image.new("RGB", (3000, 2100), "#edf2f5")
    draw = ImageDraw.Draw(image)
    draw.rectangle((80, 80, 2920, 2020), outline="#78909c", width=8)
    title = {
        "location-context": "LOCATION / CONTEXT",
        "construction-scene": "CONSTRUCTION SCENE",
        "section-or-exploded": "SECTION / EXPLODED VIEW",
        "parameter-control": "ENGINEERING VISUAL",
        "quality-safety": "QUALITY / SAFETY",
        "finished-product": "FINISHED PRODUCT VIEW",
    }.get(page.visual_type, "ENGINEERING VISUAL")
    draw.text((160, 150), title, font=ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 96), fill="#31566d")
    # 确定性离线示意，不冒充真实生图。
    if page.visual_type == "construction-scene":
        draw.polygon([(260, 1450), (1000, 850), (2200, 850), (2780, 1450)], fill="#c9b79c", outline="#6d5b4a")
        draw.polygon([(850, 1450), (1250, 1050), (1950, 1050), (2350, 1450)], fill="#7f684e")
        draw.rectangle((350, 550, 900, 820), fill="#f2b84b", outline="#785b20", width=8)
        draw.ellipse((420, 760, 590, 930), fill="#37474f"); draw.ellipse((700, 760, 870, 930), fill="#37474f")
    elif page.visual_type == "section-or-exploded":
        colors = ["#d9c4a5", "#b0bec5", "#90a4ae", "#78909c"]
        for i, color in enumerate(colors):
            y = 600 + i * 260
            draw.polygon([(500, y), (2200, y), (2580, y+140), (880, y+140)], fill=color, outline="#455a64")
    elif page.visual_type == "location-context":
        draw.line((300, 1700, 2700, 400), fill="#607d8b", width=140)
        draw.line((400, 1850, 2800, 550), fill="#f5f5f5", width=18)
        for x, y in [(500,500),(2200,1300),(1600,350)]: draw.rectangle((x,y,x+400,y+260), fill="#b0bec5", outline="#607d8b", width=6)
    elif page.visual_type == "finished-product":
        draw.rectangle((420, 620, 2580, 1580), fill="#d7d9d8", outline="#53636c", width=12)
        draw.rectangle((650, 780, 2350, 1450), fill="#89969d", outline="#455a64", width=10)
        for x in range(500, 2550, 380):
            draw.rectangle((x, 480, x+320, 650), fill="#9aa4aa", outline="#455a64", width=5)
    else:
        for i in range(5):
            draw.rounded_rectangle((250+i*500, 650, 630+i*500, 1350), radius=30, fill="#cfd8dc", outline="#607d8b", width=6)
    draw.text((160, 1860), "OFFLINE PLACEHOLDER — REPLACE WITH IMAGE API OUTPUT", font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 44), fill="#b54a42")
    return image


def parameter_labels(claims: list[Claim], limit: int = 4) -> list[str]:
    labels: list[str] = []
    for claim in claims:
        compound = r"\d+(?:\.\d+)?\s*(?:mm|cm|m)?(?:\s*[×x]\s*\d+(?:\.\d+)?\s*(?:mm|cm|m)?){1,2}"
        ranges = r"\d+(?:\.\d+)?\s*[～~:]\s*\d+(?:\.\d+)?\s*(?:mm|cm|m|d|h|MPa|%|天)?"
        scalar = r"\d+(?:\.\d+)?\s*(?:mm|cm|m|d|h|MPa|%|天|倍)"
        matches = re.findall(f"(?:{compound}|{ranges}|{scalar})", claim.text, re.I)
        for match in matches:
            value = re.sub(r"\s+", "", match)
            if value and value not in labels:
                labels.append(value)
            if len(labels) >= limit:
                return labels
    return labels


def draw_dimension(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], label: str, color: str = "#d46a2d") -> None:
    draw.line((*start, *end), fill=color, width=7)
    sx, sy = start; ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        draw.polygon([(sx, sy), (sx+34, sy-18), (sx+34, sy+18)], fill=color)
        draw.polygon([(ex, ey), (ex-34, ey-18), (ex-34, ey+18)], fill=color)
        box = draw.textbbox((0, 0), label, font=pil_font(34, True))
        tx, ty = (sx+ex-box[2])/2, sy-58
    else:
        draw.polygon([(sx, sy), (sx-18, sy+34), (sx+18, sy+34)], fill=color)
        draw.polygon([(ex, ey), (ex-18, ey-34), (ex+18, ey-34)], fill=color)
        box = draw.textbbox((0, 0), label, font=pil_font(34, True))
        tx, ty = sx+28, (sy+ey-box[3])/2
    draw.rounded_rectangle((tx-12, ty-5, tx+box[2]+12, ty+box[3]+8), radius=8, fill="white")
    draw.text((tx, ty), label, font=pil_font(34, True), fill=color)


def ocr_unexpected_text(image_path: Path) -> list[str] | None:
    if not OCR_AVAILABLE:
        return None
    try:
        import pytesseract
        text = pytesseract.image_to_string(Image.open(image_path), lang="eng+chi_sim")
    except Exception:
        return None
    return [line.strip() for line in text.splitlines() if len(line.strip()) >= 3]


def wrap_text(text: str, width: int) -> list[str]:
    # 中文按字符近似换行，ASCII 连续词亦可接受。
    return textwrap.wrap(text, width=width, break_long_words=True, break_on_hyphens=False) or [""]


def _flatten(image: Image.Image) -> Image.Image:
    """RGBA/LA/P 模式直接 convert('RGB') 会把透明区变黑，必须先合成到白底。"""
    if image.mode in ("RGBA", "LA", "P"):
        rgba = image.convert("RGBA")
        bg = Image.new("RGB", rgba.size, "white")
        bg.paste(rgba, mask=rgba.split()[-1])
        return bg
    return image.convert("RGB")


def fit_contain(image: Image.Image, size: tuple[int, int]) -> Image.Image:
    src = _flatten(image)
    ratio = min(size[0] / src.width, size[1] / src.height)
    resized = src.resize((round(src.width * ratio), round(src.height * ratio)), Image.Resampling.LANCZOS)
    result = Image.new("RGB", size, "#edf2f5")
    result.paste(resized, ((size[0] - resized.width) // 2, (size[1] - resized.height) // 2))
    return result


def render_page(page: PageSpec, claims_by_id: dict[str, Claim], visual: Image.Image, output: Path, draft: bool = True) -> Path:
    W, H = A3_PX
    image = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(image)
    navy, blue, pale, red, gray = "#17354d", "#2e6485", "#eaf0f4", "#b33939", "#4d5b64"
    draw.rectangle((0, 0, W, 260), fill=navy)
    draw.text((150, 65), page.title, font=pil_font(92, True), fill="white")
    draw.text((W-760, 95), page.page_id, font=pil_font(44), fill="#c7dbe8")
    visual_box = (150, 370, 3190, 2870)
    fitted = fit_contain(visual, (visual_box[2]-visual_box[0], visual_box[3]-visual_box[1]))
    image.paste(fitted, visual_box[:2])
    draw.rectangle(visual_box, outline="#8ba0ad", width=6)
    page_claims = [claims_by_id[cid] for cid in page.claim_ids if cid in claims_by_id]
    labels = parameter_labels(page_claims)
    if labels and page.visual_type in {"section-or-exploded", "finished-product", "construction-scene"}:
        for idx, label in enumerate(labels):
            y = visual_box[1] + 260 + idx * 250
            draw_dimension(draw, (visual_box[0]+180, y), (visual_box[0]+1000, y), label)
    legend_y = visual_box[3] - 165
    for idx, (color, text) in enumerate([(blue, "工程对象"), ("#d46a2d", "尺寸/风险标注"), ("#6f7d62", "既有环境")]):
        lx = visual_box[0] + 55 + idx * 510
        draw.rectangle((lx, legend_y, lx+56, legend_y+56), fill=color, outline="white", width=2)
        draw.text((lx+75, legend_y+5), text, font=pil_font(32, True), fill=navy)
    x0, y0, x1, y1 = 3300, 370, 4810, 2870
    draw.rounded_rectangle((x0, y0, x1, y1), radius=28, fill=pale, outline="#8ba0ad", width=5)
    draw.rectangle((x0, y0, x1, y0+150), fill=blue)
    draw.text((x0+45, y0+35), "关键控制与参数", font=pil_font(54, True), fill="white")
    claims = page_claims
    row_y = y0 + 185
    for idx, claim in enumerate(claims[:10], 1):
        lines = wrap_text(claim.text, 24)[:3]
        height = max(150, 60 + len(lines) * 50)
        if row_y + height > y1 - 40: break
        fill = "#fff0db" if claim.status in BLOCKING_STATUS else "white"
        draw.rounded_rectangle((x0+30, row_y, x1-30, row_y+height-14), radius=18, fill=fill, outline="#b5c2ca", width=3)
        draw.ellipse((x0+55, row_y+35, x0+115, row_y+95), fill=red if claim.claim_type in {"PROHIBITION", "RISK_CONTROL"} else blue)
        draw.text((x0+74, row_y+42), str(idx), font=pil_font(30, True), fill="white")
        ty = row_y + 24
        for line in lines:
            draw.text((x0+140, ty), line, font=pil_font(38), fill=navy)
            ty += 50
        draw.text((x1-310, row_y+height-58), claim.claim_id, font=pil_font(25), fill=gray)
        row_y += height
    label = VISUAL_LABELS.get(page.truth_level, VISUAL_LABELS["generic"])
    if page.truth_level == "generic":
        disclaimer = "AI施工示意图｜非现场实景｜非按比例｜不作为施工放样及尺寸验收依据"
    else:
        disclaimer = f"{label}｜位置及尺寸以经批准图纸和现场测量成果为准"
    draw.rectangle((0, 3030, W, H), fill="#f3f5f6")
    draw.text((150, 3075), disclaimer, font=pil_font(40, True), fill=red)
    source_ids = sorted({claims_by_id[cid].source_id for cid in page.claim_ids if cid in claims_by_id})
    draw.text((150, 3165), "来源：" + "、".join(source_ids or ["无"]), font=pil_font(34), fill=gray)
    draw.text((150, 3240), "事实索引：" + "、".join(page.claim_ids[:14]), font=pil_font(30), fill=gray)
    if page.scoring_points:
        draw.text((3300, 3075), "评分响应：" + "；".join(page.scoring_points), font=pil_font(34, True), fill=navy)
    if draft:
        draw.text((W-780, H-160), "DRAFT / 待总工审核", font=pil_font(48, True), fill=red)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, dpi=(300, 300))
    return output


def register_pdf_font() -> str:
    name = "BidVizCN"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, font_path(False)))
    return name


def compile_pdf(page_images: list[Path], output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(output), pagesize=landscape(A3), pageCompression=1)
    width, height = landscape(A3)
    for page in page_images:
        c.drawImage(str(page), 0, 0, width=width, height=height, preserveAspectRatio=False)
        c.showPage()
    c.save()
    return output


def audit(project: Project, page_paths: list[Path] | None = None, pdf_path: Path | None = None, final: bool = False) -> dict[str, Any]:
    findings: list[dict[str, Any]] = []
    claims = {c.claim_id: c for c in project.claims}
    for claim in project.claims:
        try: claim.validate()
        except ValueError as exc: findings.append({"severity": "error", "code": "invalid-claim", "message": str(exc)})
    for page in project.pages:
        missing = [cid for cid in page.claim_ids if cid not in claims]
        if missing: findings.append({"severity": "error", "code": "missing-claim", "page": page.page_id, "message": ", ".join(missing)})
        if len(page.reference_assets) > MAX_REFERENCES:
            findings.append({"severity": "error", "code": "too-many-references", "page": page.page_id, "message": str(len(page.reference_assets))})
    unresolved = [i for i in project.issues if i.blocking and not i.resolved]
    if final and unresolved:
        findings.append({"severity": "error", "code": "blocking-issues", "message": f"{len(unresolved)} 个阻断问题未关闭"})
    if final:
        unsafe = [c.claim_id for c in project.claims if c.status in BLOCKING_STATUS]
        if unsafe:
            findings.append({"severity": "error", "code": "unapproved-claims", "message": "、".join(unsafe)})
    for path in page_paths or []:
        try:
            with Image.open(path) as im:
                if im.size != A3_PX:
                    findings.append({"severity": "error", "code": "wrong-pixel-size", "page": path.name, "message": str(im.size)})
                dpi = im.info.get("dpi", (0, 0))
                if min(dpi or (0, 0)) < 299:
                    findings.append({"severity": "warning", "code": "low-dpi", "page": path.name, "message": str(dpi)})
        except Exception as exc:
            findings.append({"severity": "error", "code": "invalid-image", "page": path.name, "message": str(exc)})
    visual_paths: list[Path] = getattr(project, "_visual_paths", [])
    if visual_paths:
        if not OCR_AVAILABLE:
            findings.append({"severity": "warning", "code": "ocr-unavailable", "message": "未安装可用 OCR，在线主视觉伪文字须人工复核"})
        else:
            for path in visual_paths:
                unexpected = ocr_unexpected_text(path)
                if unexpected:
                    findings.append({"severity": "error" if final else "warning", "code": "generated-text", "page": path.name, "message": "；".join(unexpected[:5])})
    pdf_pages = None
    pdf_size = None
    if pdf_path and pdf_path.exists() and fitz:
        doc = fitz.open(pdf_path)
        pdf_pages = len(doc)
        if doc:
            rect = doc[0].rect
            pdf_size = [round(rect.width, 2), round(rect.height, 2)]
            expected = landscape(A3)
            if abs(rect.width - expected[0]) > 1 or abs(rect.height - expected[1]) > 1:
                findings.append({"severity": "error", "code": "wrong-pdf-size", "message": str(pdf_size)})
        if page_paths is not None and pdf_pages != len(page_paths):
            findings.append({"severity": "error", "code": "page-count-mismatch", "message": f"PDF={pdf_pages}, PNG={len(page_paths)}"})
    return {
        "generated_at": utc_now(),
        "mode": "final" if final else "draft",
        "passed": not any(f["severity"] == "error" for f in findings),
        "summary": {"claims": len(project.claims), "pages": len(project.pages), "issues": len(project.issues), "unresolved_blockers": len(unresolved), "pdf_pages": pdf_pages, "pdf_size_points": pdf_size},
        "findings": findings,
    }


def project_to_dict(project: Project) -> dict[str, Any]:
    return {
        "name": project.name,
        "sources": [asdict(x) for x in project.sources],
        "claims": [asdict(x) for x in project.claims],
        "issues": [asdict(x) for x in project.issues],
        "pages": [asdict(x) for x in project.pages],
    }


def project_from_dict(data: dict[str, Any]) -> Project:
    return Project(
        name=data["name"],
        sources=[Source(**x) for x in data.get("sources", [])],
        claims=[Claim(**x) for x in data.get("claims", [])],
        issues=[Issue(**x) for x in data.get("issues", [])],
        pages=[PageSpec(**x) for x in data.get("pages", [])],
    )


def write_trace_csv(project: Project, path: Path) -> None:
    import csv
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["claim_id", "work_item_id", "type", "status", "text", "source_id", "locator", "section_path"])
        for c in project.claims:
            writer.writerow([c.claim_id, c.work_item_id, c.claim_type, c.status, c.text, c.source_id, c.locator, c.section_path])


def write_review(project: Project, report: dict[str, Any], path: Path) -> None:
    lines = [f"# 审查报告：{project.name}", "", f"- 生成时间：{report['generated_at']}", f"- 模式：{report['mode']}", f"- 自动门禁：{'通过' if report['passed'] else '未通过'}", f"- 事实：{len(project.claims)}", f"- 图页：{len(project.pages)}", f"- 未关闭阻断项：{report['summary']['unresolved_blockers']}", "", "## 自动检查发现", ""]
    if report["findings"]:
        lines += [f"- [{x['severity']}] `{x['code']}`：{x['message']}" for x in report["findings"]]
    else:
        lines.append("- 无自动检查错误。")
    lines += ["", "## 总工人工复核", "", "- [ ] 项目位置与参考图一致", "- [ ] 工程构造和施工机械关系合理", "- [ ] 所有关键参数与原文或批准值一致", "- [ ] AI 通用图未冒充现场实景", "- [ ] 评分点响应准确且无夸大", "- [ ] 批准终稿导出"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def run_project(project: Project, root: Path, online: bool = False, final: bool = False) -> dict[str, Any]:
    work = root / "work"
    out = root / "output"
    visuals = work / "visuals"
    page_dir = out / "pages"
    claims_by_id = {c.claim_id: c for c in project.claims}
    for page in project.pages:
        page.prompt = compile_prompt(page, claims_by_id)
    save_json(work / "project.json", project_to_dict(project))
    page_paths: list[Path] = []
    api = ImageAPI(provider_name=os.getenv("CUSTOM_IMAGE_PROVIDER")) if online else None
    generation_records: list[dict[str, Any]] = []
    for page in project.pages:
        visual_path = visuals / f"{page.page_id}.png"
        references = [root / "inputs" / p for p in page.reference_assets]
        references = [p for p in references if p.exists()]
        if online:
            assert api is not None
            record: dict[str, Any] = {"page_id": page.page_id}
            api.generate(page.prompt, visual_path, references, recorder=record)
            generation_records.append(record)
            visual = Image.open(visual_path).convert("RGB")
        else:
            visual = placeholder_visual(page)
            visual_path.parent.mkdir(parents=True, exist_ok=True)
            visual.save(visual_path)
        page_path = page_dir / f"{page.page_id}.png"
        render_page(page, claims_by_id, visual, page_path, draft=not final)
        page_paths.append(page_path)
    project._visual_paths = sorted(visuals.glob("*.png")) if online else []
    pdf_path = out / ("施工组织设计可视化图册-A3-FINAL.pdf" if final else "施工组织设计可视化图册-A3-DRAFT.pdf")
    compile_pdf(page_paths, pdf_path)
    write_trace_csv(project, out / "source-trace.csv")
    if generation_records:
        save_json(work / "generation-log.json", generation_records)
    report = audit(project, page_paths, pdf_path, final=final)
    save_json(out / "audit-report.json", report)
    write_review(project, report, out / "review-report.md")
    if final and not report["passed"]:
        pdf_path.unlink(missing_ok=True)
        raise BidVizError("终稿门禁未通过，已阻止 FINAL PDF 输出；详见 audit-report.json")
    return {"pdf": str(pdf_path), "pages": [str(x) for x in page_paths], "report": report}


def cmd_probe(args: argparse.Namespace) -> None:
    api = ImageAPI(provider_name=args.provider, model=args.model)
    models = api.models()
    print(json.dumps({"endpoint": api.endpoint, "model": api.model, "model_available": api.model in models, "models": models}, ensure_ascii=False))


def cmd_manifest(args: argparse.Namespace) -> None:
    save_json(Path(args.output), build_manifest(Path(args.input)))
    print(args.output)


def cmd_extract(args: argparse.Namespace) -> None:
    source_path = Path(args.source)
    claims = extract_claims(extract_lines(source_path), args.source_id, args.work_item)
    issues = detect_issues(claims)
    save_json(Path(args.output), {"claims": [asdict(c) for c in claims], "issues": [asdict(i) for i in issues]})
    print(f"claims={len(claims)} issues={len(issues)}")


def cmd_render(args: argparse.Namespace) -> None:
    project = project_from_dict(load_json(Path(args.project)))
    result = run_project(project, Path(args.root), online=args.online, final=args.final)
    print(result["pdf"])


def cmd_audit(args: argparse.Namespace) -> None:
    project = project_from_dict(load_json(Path(args.project)))
    pages = sorted(Path(args.pages).glob("*.png")) if args.pages else []
    project._visual_paths = sorted(Path(args.visuals).glob("*.png")) if args.visuals else []
    report = audit(project, pages, Path(args.pdf) if args.pdf else None, final=args.final)
    if args.output: save_json(Path(args.output), report)
    print(json.dumps(report["summary"], ensure_ascii=False))
    raise SystemExit(0 if report["passed"] else 1)


def cmd_demo(args: argparse.Namespace) -> None:
    if args.provider:
        os.environ["CUSTOM_IMAGE_PROVIDER"] = args.provider
    if args.model:
        os.environ["CUSTOM_IMAGE_MODEL"] = args.model
    root = Path(args.root)
    data = load_json(root / "examples" / "poc" / "project.json")
    project = project_from_dict(data)
    result = run_project(project, root / "examples" / "poc", online=not args.offline, final=args.final)
    print(result["pdf"])


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    x = sub.add_parser("probe"); x.add_argument("--provider"); x.add_argument("--model"); x.set_defaults(func=cmd_probe)
    x = sub.add_parser("manifest"); x.add_argument("input"); x.add_argument("-o", "--output", default="work/manifest.json"); x.set_defaults(func=cmd_manifest)
    x = sub.add_parser("extract"); x.add_argument("source"); x.add_argument("--source-id", required=True); x.add_argument("--work-item", default="AUTO"); x.add_argument("-o", "--output", default="work/extracted.json"); x.set_defaults(func=cmd_extract)
    x = sub.add_parser("render"); x.add_argument("project"); x.add_argument("--root", default="."); x.add_argument("--online", action="store_true"); x.add_argument("--final", action="store_true"); x.set_defaults(func=cmd_render)
    x = sub.add_parser("audit"); x.add_argument("project"); x.add_argument("--pages"); x.add_argument("--visuals"); x.add_argument("--pdf"); x.add_argument("--output"); x.add_argument("--final", action="store_true"); x.set_defaults(func=cmd_audit)
    x = sub.add_parser("demo"); x.add_argument("--root", default=str(Path(__file__).parent)); x.add_argument("--offline", action="store_true"); x.add_argument("--final", action="store_true"); x.add_argument("--provider"); x.add_argument("--model"); x.set_defaults(func=cmd_demo)
    return p


def main() -> None:
    args = parser().parse_args()
    try:
        args.func(args)
    except BidVizError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
